import sys, copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

PX = 6350  # EMU per CSS px on a 1920-wide slide
def E(v): return Emu(int(round(v * PX)))
def C(h): return RGBColor.from_string(h.lstrip('#').upper())

BG='0b0908'; CREAM='efe9e2'; TEXT='e3dcd3'; MUTED='a99f93'; COPPER='d9a070'; COPPER2='e9bd93'
CARD='141210'; CARDLINE='262321'; CHIP='2a1f17'; CHIPLINE='4a3626'; CHIPTXT='f3cfae'
FONT='Inter'

prs = Presentation()
prs.slide_width = E(1920); prs.slide_height = E(1080)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = C(BG)
sh = s.shapes

def rect(x,y,w,h,fill,line=None,radius=None,dash=False,lw=1):
    shp = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, E(x),E(y),E(w),E(h))
    if radius: shp.adjustments[0] = radius/min(w,h)
    if fill: shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    else: shp.fill.background()
    if line:
        shp.line.color.rgb = C(line); shp.line.width = E(lw)
        if dash:
            ln = shp.line._get_or_add_ln(); d = etree.SubElement(ln, qn('a:prstDash')); d.set('val','dash')
    else: shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ''
    return shp

def tb(x,y,w,h,anchor=MSO_ANCHOR.TOP):
    t = sh.add_textbox(E(x),E(y),E(w),E(h)); tf = t.text_frame
    tf.word_wrap = True; tf.auto_size = None
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf, m, 0)
    tf.vertical_anchor = anchor
    return tf

def run(p, text, px, color, bold=False, spc=None, u=False):
    r = p.add_run(); r.text = text; f = r.font
    f.name = FONT; f.size = Pt(px*0.5); f.bold = bold; f.color.rgb = C(color); f.underline = u
    if spc: r._r.get_or_add_rPr().set('spc', str(spc))
    return r

def lnspc(p, pct):
    pPr = p._p.get_or_add_pPr()
    l = etree.SubElement(pPr, qn('a:lnSpc')); etree.SubElement(l, qn('a:spcPct')).set('val', str(int(pct*1000)))
    return pPr

def bullets(tf, items, px=19.5):
    first = True
    for parts in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        pPr = lnspc(p, 112)
        sa = etree.SubElement(pPr, qn('a:spcAft')); etree.SubElement(sa, qn('a:spcPts')).set('val', '500')
        pPr.set('marL', str(22*PX)); pPr.set('indent', str(-22*PX))
        bc = etree.SubElement(pPr, qn('a:buClr')); etree.SubElement(bc, qn('a:srgbClr')).set('val', COPPER.upper())
        etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
        etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
        for txt, b in parts:
            run(p, txt, px, 'ffffff' if b else TEXT, bold=b)

def label(x,y,text):
    tf = tb(x,y,700,22); run(tf.paragraphs[0], text.upper(), 16, COPPER, bold=True, spc=200)

# ---------- owl logo (vector freeforms) ----------
def bez(p0,p1,p2,p3,n=14):
    return [((1-t)**3*p0[0]+3*(1-t)**2*t*p1[0]+3*(1-t)*t**2*p2[0]+t**3*p3[0],
             (1-t)**3*p0[1]+3*(1-t)**2*t*p1[1]+3*(1-t)*t**2*p2[1]+t**3*p3[1]) for t in [i/n for i in range(1,n+1)]]
def free(ox,oy,k,pts,fill,line=None,lw=0):
    fb = sh.build_freeform(E(ox+pts[0][0]*k), E(oy+pts[0][1]*k), scale=1.0)
    fb.add_line_segments([(E(ox+x*k),E(oy+y*k)) for x,y in pts[1:]], close=True)
    shp = fb.convert_to_shape(); shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
    if line: shp.line.color.rgb = C(line); shp.line.width = E(lw)
    else: shp.line.fill.background()
    shp.shadow.inherit=False
    return shp
def oval(ox,oy,k,cx,cy,r,fill):
    shp = sh.add_shape(MSO_SHAPE.OVAL, E(ox+(cx-r)*k),E(oy+(cy-r)*k),E(2*r*k),E(2*r*k))
    shp.fill.solid(); shp.fill.fore_color.rgb = C(fill); shp.line.fill.background(); shp.shadow.inherit=False
def owl(ox,oy,size):
    k = size/256
    body = [(48,120)] + bez((48,120),(48,186),(96,230),(128,244)) + bez((128,244),(160,230),(208,186),(208,120))
    free(ox,oy,k,body,CREAM)
    for cx in (92,164): oval(ox,oy,k,cx,160,30,BG)
    for cx in (96,160): oval(ox,oy,k,cx,164,13,CREAM)
    free(ox,oy,k,[(128,192),(137,203),(128,216),(119,203)],BG)
    cap = [(32,118)] + bez((32,118),(32,66),(76,38),(128,38)) + bez((128,38),(180,38),(224,66),(224,118)) \
        + bez((224,118),(198,112),(150,112),(128,146)) + bez((128,146),(106,112),(58,112),(32,118))
    free(ox,oy,k,cap,CREAM,BG,10*k)
    st = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(ox+119*k),E(oy+16*k),E(18*k),E(34*k))
    st.adjustments[0]=0.5; st.fill.solid(); st.fill.fore_color.rgb=C(CREAM); st.line.fill.background(); st.shadow.inherit=False

# ---------- header ----------
owl(64,44,74)
tf = tb(156,44,600,74,MSO_ANCHOR.MIDDLE); run(tf.paragraphs[0],'PreDiligence',44,CREAM,bold=True,spc=-120)
tf = tb(1256,48,600,60); 
for i,t in enumerate(['Accel · Future Founder Challenge','Data-Centric Agents · Paris']):
    p = tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment = PP_ALIGN.RIGHT; lnspc(p,115); run(p,t,17,MUTED)

tf = tb(64,138,1600,120); p = tf.paragraphs[0]; lnspc(p,98)
run(p,'The AI analyst that turns a startup name into a ',50,CREAM,bold=True,spc=-180)
run(p,'sourced pre-call memo in 2 minutes',50,COPPER2,bold=True,spc=-180)
run(p,', and flags what the databases got wrong.',50,CREAM,bold=True,spc=-180)

# ---------- row 1 ----------
cols = [
 ("Who it's for", [[('Analysts and partners at pre-seed and seed funds',True),(', who screen hundreds of startups a month.',False)],
                   [('Each first look takes ',False),('2 to 4 hours',True),(' of research on public data that is scattered, stale and contradictory.',False)],
                   [('Incumbents serve PE, not seed: DiligenceSquared charges ',False),('~$50k per project',True),('.',False)]]),
 ("Why now", [[('Live web search built for agents',True),(' is now cheap and fast (Keenable).',False)],
              [('LLMs can finally ',False),('extract and cross-check facts',True),(', not just summarise.',False)],
              [('The market pays: Rogo passed ',False),('$50M ARR',True),(', and Accel co-led ',False),("Bridgetown Research's $19M Series A",True),('. Nobody serves early stage.',False)]]),
 ("First users", [[('The VCs in this room',True),(': free pilot from tomorrow, on their real deal flow.',False)],
                  [('Paris seed funds, scout programs and angel networks, through warm intros.',False)],
                  [('A weekly public post: ',False),('"what the databases got wrong this week"',True),('. Then €200–500 per seat per month.',False)]]),
]
for i,(lab,items) in enumerate(cols):
    x = 64 + i*606
    rect(x,272,580,285,CARD,CARDLINE,radius=22)
    label(x+28,294,lab)
    bullets(tb(x+28,326,526,224), items)

# ---------- row 2 left: how it works ----------
rect(64,580,945,456,CARD,CARDLINE,radius=22)
label(92,602,'How it works · data at the core')
def chip(x,y,w,text,src=False):
    shp = rect(x,y,w,44,'1f1c1a' if src else CHIP,'3a3633' if src else CHIPLINE,radius=12)
    tf = shp.text_frame
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,m,0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap=False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; run(p,text,18,CREAM if src else CHIPTXT,bold=True)
def arrow(x,y):
    tf = tb(x,y,30,44,MSO_ANCHOR.MIDDLE); p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; run(p,'→',22,'6f665d')
row1 = [('Keenable live web + X',224,True),('Planner',104,False),('4 research agents',196,False),('Fact extractor',162,False)]
x = 92
for i,(t,w,src) in enumerate(row1):
    chip(x,634,w,t,src); x += w
    arrow(x+2,634); x += 34
x = 92
for i,(t,w) in enumerate([('Verifier',104),('Memo + 5 call questions',250)]):
    chip(x,690,w,t); x += w
    if i==0: arrow(x+2,690); x += 34
rect(92,754,889,130,'0f0d0b',CHIPLINE,radius=14,dash=True)
tf = tb(120,772,840,100); p = tf.paragraphs[0]; lnspc(p,118)
run(p,'Aiven PostgreSQL is the memory. ',20,'ffffff',bold=True)
run(p,'Every fact is stored with its source, date and confidence. No source, no fact. The primary source beats stale databases, and every re-check shows exactly what changed.',20,TEXT)

# ---------- row 2 right: proof ----------
rect(1035,580,821,456,CARD,CARDLINE,radius=22)
label(1063,602,'Proof')
for (bx,fill,line,big,small,sc) in [(1063,'1d1a18',None,'"Active"','YC directory, PitchBook and Tracxn on Menza (YC W26)',MUTED),
                                    (1451,'33261c','6b4e36','"Acquired"',"PreDiligence found it on Menza's own blog, with the source and date",'f6d8bc')]:
    rect(bx,634,377,116,fill,line,radius=14)
    tf = tb(bx+18,646,343,96); p=tf.paragraphs[0]; run(p,big,24,'ffffff',bold=True)
    p=tf.add_paragraph(); lnspc(p,112); run(p,small,18,sc)
cx = [1063,1380,1611]
tf = tb(cx[0],774,300,20); run(tf.paragraphs[0],'10 YC F26 STARTUPS',15,MUTED,bold=True,spc=100)
tf = tb(cx[1],774,220,20); run(tf.paragraphs[0],'LLM ALONE',15,MUTED,bold=True,spc=100)
tf = tb(cx[2],774,220,20); run(tf.paragraphs[0],'PREDILIGENCE',15,MUTED,bold=True,spc=100)
rows = [('Correct facts','  __%  ','  __%  '),('Facts with a source','  __%  ','  __%  '),('Time per startup','2–4 h by hand','  __ min  ')]
for i,(a,b,c) in enumerate(rows):
    y = 806 + i*42
    ln = sh.add_connector(MSO_CONNECTOR.STRAIGHT, E(1063),E(y),E(1828),E(y)); ln.line.color.rgb=C(CARDLINE); ln.line.width=E(1)
    tf = tb(cx[0],y+9,300,26); run(tf.paragraphs[0],a,20,TEXT)
    tf = tb(cx[1],y+9,220,26); run(tf.paragraphs[0],b,20,COPPER if '__' in b else TEXT,u='__' in b)
    tf = tb(cx[2],y+9,220,26); run(tf.paragraphs[0],c,20,COPPER,bold=True,u=True)

out = sys.argv[1]
prs.save(out)
